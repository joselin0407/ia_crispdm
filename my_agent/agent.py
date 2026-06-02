# my_agent/agent.py
import os
import json
import nbformat
import docx
from docx.shared import Pt
from google.adk.agents.llm_agent import Agent

# ==========================================
# 1. DEFINICIÓN DE HERRAMIENTAS (TOOLS)
# ==========================================

def extract_notebook_metadata(filepath: str) -> str:
    """
    Lee un archivo Jupyter Notebook (.ipynb) de la ruta local, extrae su contenido 
    (celdas de código, celdas de markdown, librerías importadas) y lo retorna como un JSON estructurado.
    
    Use esta herramienta ÚNICAMENTE cuando el usuario proporcione la ruta o nombre de un archivo .ipynb para analizar.
    
    Args:
        filepath: Ruta del archivo .ipynb en el disco local.
    """
    if not os.path.exists(filepath):
        return f"Error: El archivo en la ruta '{filepath}' no existe en este directorio."
        
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            nb = nbformat.read(f, as_version=4)
        
        extracted_data = {
            "notebook_name": os.path.basename(filepath),
            "total_cells": len(nb.cells),
            "markdown_blocks": [],
            "code_blocks": [],
            "libraries_used": set()
        }

        for idx, cell in enumerate(nb.cells):
            if cell.cell_type == 'markdown':
                extracted_data["markdown_blocks"].append({
                    "cell_index": idx,
                    "content": cell.source
                })
            elif cell.cell_type == 'code':
                code_content = cell.source
                extracted_data["code_blocks"].append({
                    "cell_index": idx,
                    "code": code_content
                })
                # Identifica imports básicos en el notebook
                for line in code_content.split('\n'):
                    if line.strip().startswith(('import ', 'from ')):
                        extracted_data["libraries_used"].add(line.strip())

        extracted_data["libraries_used"] = list(extracted_data["libraries_used"])
        return json.dumps(extracted_data, indent=4, ensure_ascii=False)
        
    except Exception as e:
        return f"Error técnico al leer el notebook: {str(e)}"


def save_markdown_as_docx(markdown_content: str, filename: str = "documento_CRISP_DM.docx") -> str:
    """
    Convierte un texto con formato Markdown en un archivo de Microsoft Word (.docx) formal 
    y lo guarda en el disco local en el directorio actual.
    
    Use esta herramienta cuando termine de redactar la documentación CRISP-DM y deba exportar 
    el resultado a un archivo de Word (.docx) para entregárselo al usuario.
    
    Args:
        markdown_content: El texto completo de la documentación generada por la IA en formato Markdown.
        filename: El nombre del archivo .docx que se va a guardar.
    """
    try:
        doc = docx.Document()
        
        # Configuración de estilos estéticos base
        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)

        doc.add_heading('Documentación Oficial de Proyecto ML (CRISP-DM)', level=0)
        
        lines = markdown_content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # Formateo simple de encabezados de Markdown a Word
            if line.startswith('### '):
                doc.add_heading(line.replace('### ', ''), level=3)
            elif line.startswith('## '):
                doc.add_heading(line.replace('## ', ''), level=2)
            elif line.startswith('# '):
                doc.add_heading(line.replace('# ', ''), level=1)
            elif line.startswith('- '):
                doc.add_paragraph(line.replace('- ', ''), style='List Bullet')
            else:
                doc.add_paragraph(line)
                
        doc.save(filename)
        return f"¡Éxito! El documento de Word ha sido autogenerado y guardado como '{filename}' en tu carpeta de proyecto."
        
    except Exception as e:
        return f"Error técnico al generar el documento Word: {str(e)}"

CRISP_DM_INSTRUCTIONS = """
Eres AIDA (Agente Inteligente de Documentación Analítica), un experto técnico corporativo especializado en Ciencia de Datos, Gobierno de Datos y la metodología CRISP-DM. 

Tu objetivo principal es automatizar la documentación de proyectos de Machine Learning, analizando el contexto (código Python, comentarios, celdas de Markdown y metadatos) extraído de Jupyter Notebooks y generando documentos oficiales estandarizados.

INSTRUCCIONES DE OPERACIÓN (Flujo de Trabajo):
1. EXTRACCIÓN: Si el usuario te da el nombre o la ruta de un archivo '.ipynb', invoca inmediatamente la herramienta 'extract_notebook_metadata' pasándole esa ruta para leer su contenido.
2. ANÁLISIS: Analiza rigurosamente el JSON retornado por la herramienta. Utiliza únicamente ese contexto para estructurar el documento oficial.
3. REDACCIÓN: Escribe la documentación de forma completa, analítica y técnica en formato Markdown, siguiendo estrictamente la estructura CRISP-DM detallada más abajo.
4. EXPORTACIÓN: Una vez que hayas terminado de redactar toda la documentación en Markdown, invoca de forma autónoma la herramienta 'save_markdown_as_docx' para exportar tu redacción a un archivo Word (.docx).
5. NOTIFICACIÓN: Finalmente, dile al usuario que el archivo de Word está listo para su descarga y ofrécele un muy breve resumen ejecutivo de los hallazgos del modelo.

REGLAS DE GENERACIÓN OBLIGATORIAS:
- Tono y Estilo: Formal, técnico, objetivo y corporativo (apropiado para la Dirección de Data y Analítica).
- Formato: Usa encabezados (##, ###), viñetas y tablas de Markdown para organizar la información de forma visualmente atractiva y fácil de leer.
- CERO ALUCINACIONES (Crítico): Basarás tu documentación ÚNICAMENTE en la información extraída del notebook. NO inventes datos, métricas, librerías o variables. Si el código no evidencia información sobre una fase específica (por ejemplo, no hay pasos de despliegue), DEBES escribir explícitamente: "No se evidencia información sobre esta fase en el código analizado."

ESTRUCTURA OBLIGATORIA DEL DOCUMENTO (Metodología CRISP-DM):
Tu redacción debe contener siempre estas 6 secciones principales:

## 1. Comprensión del Negocio (Business Understanding)
- Identifica (a partir de los comentarios, el nombre del archivo o variables target) el objetivo del modelo y el problema de negocio que intenta resolver.

## 2. Comprensión de los Datos (Data Understanding)
- Enumera las fuentes de datos utilizadas (archivos leídos como CSV, Parquet, o bases de datos conectadas).
- Describe el volumen de datos (shape), las principales variables identificadas y cualquier análisis exploratorio inicial (EDA) que se note en el código (ej. gráficos generados, validación de nulos o duplicados).

## 3. Preparación de los Datos (Data Preparation)
- Detalla los pasos técnicos realizados sobre el set de datos: limpieza, imputación de nulos, transformación de variables (encoding, escalado), y selección de características (feature engineering). 
- Menciona explícitamente las librerías de Python utilizadas para este propósito (ej. Pandas, NumPy, Scikit-learn).

## 4. Modelado (Modeling)
- Especifica el tipo de modelo(s) o algoritmos entrenados (ej. Random Forest, Regresión Lineal, XGBoost, Redes Neuronales).
- Crea una tabla comparativa en Markdown con los hiperparámetros clave configurados en el código de entrenamiento.

## 5. Evaluación (Evaluation)
- Reporta las métricas de rendimiento extraídas (ej. Accuracy, RMSE, F1-Score, matriz de confusión).
- Explica brevemente, de manera técnica, qué significa el resultado de las métricas en el contexto del modelo.

## 6. Despliegue (Deployment)
- Identifica si el notebook contiene pasos para exportar, guardar o empaquetar el modelo (ej. uso de `pickle`, `joblib`, u ONNX, o guardado en plataformas de Cloud)."""

# ==========================================
# 3. CREACIÓN DEL AGENTE REGISTRANDO LAS TOOLS
# ==========================================

root_agent = Agent(
    model='gemini-3.1-flash-lite',#'gemini-1.5-pro', #'gemini-2.5-flash', 
    name='AIDA',
    description="Agente de IA especializado en documentar proyectos ML bajo metodología CRISP-DM utilizando herramientas locales.",
    instruction=CRISP_DM_INSTRUCTIONS,
    tools=[extract_notebook_metadata, save_markdown_as_docx]  # <-- AQUÍ REGISTRAMOS TUS HERRAMIENTAS
)
# ==========================================
# MENSAJE DE BIENVENIDA AUTOMÁTICO EN CONSOLA
# ==========================================
print("\n" + "═"*70)
print(" 🤖  ¡BIENVENIDO A AIDA!  🤖 ")
print(" Agente Inteligente de Documentación Analítica bajo Metodología CRISP-DM")
print("═"*70)
print(" Estoy lista para ayudarte a automatizar y gobernar tus proyectos.")
print(" Puedes pedirme cosas como:")
print(" 👉 'Analiza el notebook modelo_clientes.ipynb y guarda el reporte'")
print(" 👉 'Genera la documentación CRISP-DM del archivo local pruebas.ipynb'")
print("═"*70 + "\n")