# my_agent/agentv2.py
import os
import nbformat
import docx
from docx.shared import Pt
from google.adk.agents import Agent

# =====================================================================
# 1. NUEVA HERRAMIENTA: ANALIZADOR DINÁMICO DE REPOSITORIOS (REPLACES JSON)
# =====================================================================

def analyze_repository_contents(repo_path: str) -> str:
    """
    Escanea de forma recursiva una carpeta o repositorio local, lee todos los archivos 
    de código fuente de Python (.py) y Jupyter Notebooks (.ipynb), y consolida sus contenidos 
    en un único bloque de texto estructurado para que el agente de IA lo analice.
    
    Args:
        repo_path: Ruta de la carpeta o repositorio local que se va a analizar.
    """
    if not os.path.exists(repo_path):
        return f"Error: La ruta de la carpeta '{repo_path}' no existe en este directorio."
        
    consolidated_code = []
    
    try:
        # Recorremos recursivamente todos los archivos de la carpeta
        for root, dirs, files in os.walk(repo_path):
            # Omitimos carpetas ocultas de Git, ADK o entornos virtuales por rendimiento
            dirs[:] = [d for d in dirs if not d.startswith('.') and d not in ['__pycache__', 'venv', 'env']]
            
            for file in files:
                filepath = os.path.join(root, file)
                
                # Procesamos archivos de código de Python (.py)
                if file.endswith('.py'):
                    try:
                        with open(filepath, 'r', encoding='utf-8') as f:
                            content = f.read()
                        consolidated_code.append(f"=== ARCHIVO FUENTE PYTHON: {file} ===\n{content}\n")
                    except Exception as e:
                        consolidated_code.append(f"=== ERROR AL LEER ARCHIVO: {file} ({str(e)}) ===\n")
                        
                # Procesamos archivos de Jupyter Notebooks (.ipynb)
                elif file.endswith('.ipynb'):
                    try:
                        with open(filepath, 'r', encoding='utf-8') as f:
                            nb = nbformat.read(f, as_version=4)
                        
                        notebook_text = []
                        for idx, cell in enumerate(nb.cells):
                            if cell.cell_type == 'markdown':
                                notebook_text.append(f"[Celda Markdown {idx}]:\n{cell.source}\n")
                            elif cell.cell_type == 'code':
                                notebook_text.append(f"[Celda Código {idx}]:\n{cell.source}\n")
                        
                        consolidated_code.append(f"=== JUPYTER NOTEBOOK: {file} ===\n" + "\n".join(notebook_text) + "\n")
                    except Exception as e:
                        consolidated_code.append(f"=== ERROR AL LEER NOTEBOOK: {file} ({str(e)}) ===\n")
                
                # Procesamos archivos de texto (.txt)
                elif file.endswith('.txt'):
                    try:
                        with open(filepath, 'r', encoding='utf-8') as f:
                            content = f.read()
                        consolidated_code.append(f"=== ARCHIVO DE TEXTO: {file} ===\n{content}\n")
                    except Exception as e:
                        consolidated_code.append(f"=== ERROR AL LEER ARCHIVO DE TEXTO: {file} ({str(e)}) ===\n")

        if not consolidated_code:
            return f"No se encontraron archivos de código (.py, .ipynb o .txt) válidos para analizar en '{repo_path}'."
            
        return "\n\n".join(consolidated_code)
        
    except Exception as e:
        return f"Error técnico durante el escaneo del repositorio: {str(e)}"


# =====================================================================
# 2. HERRAMIENTA DE EXPORTACIÓN A WORD (DOCX)
# =====================================================================

def save_markdown_as_docx(markdown_content: str, filename: str = "documento_CRISP_DM.docx") -> str:
    """
    Convierte un texto con formato Markdown en un archivo de Microsoft Word (.docx) formal 
    y lo guarda en el disco local en el directorio actual.
    
    Args:
        markdown_content: El texto completo de la documentación generada por la IA en formato Markdown.
        filename: El nombre del archivo .docx que se va a guardar.
    """
    try:
        doc = docx.Document()
        
        # Estilos corporativos base
        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)

        doc.add_heading('Documentación Oficial de Proyecto ML (CRISP-DM)', level=0)
        
        lines = markdown_content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            if line.startswith('### '):
                doc.add_heading(line.replace('### ', ''), level=3)
            elif line.startswith('## '):
                doc.add_heading(line.replace('## ', ''), level=2)
            elif line.startswith('# '):
                doc.add_heading(line.replace('# ', ''), level=1)
            elif line.startswith('- '):
                doc.add_paragraph(line.replace('- ', ''), style='List Bullet')
            else:
                doc.add_paragraph(line)
                
        doc.save(filename)
        return f"¡Éxito! El documento de Word ha sido autogenerado y guardado como '{filename}' en tu carpeta de proyecto."
        
    except Exception as e:
        return f"Error técnico al generar el documento Word: {str(e)}"


# =====================================================================
# 3. PROMPT DE INSTRUCCIÓN ENRIQUECIDO DE CRISP-DM (MÉTODO PROFUNDO POR FASES)
# =====================================================================

CRISP_DM_INSTRUCTIONS = """
Eres AIDA (Agente Inteligente de Documentación Analítica), un experto técnico senior especializado en Ciencia de Datos, Gobierno de Datos y la metodología CRISP-DM. 

Tu objetivo principal es analizar de forma profunda todo el código fuente, comentarios y notebooks de una carpeta o repositorio de proyectos de Machine Learning, para generar de manera autónoma el documento técnico oficial de gobierno.

INSTRUCCIONES DE OPERACIÓN:
1. Si el usuario te proporciona el nombre o la ruta de una carpeta o repositorio, ejecuta inmediatamente la herramienta 'analyze_repository_contents' pasándole esa ruta.
2. Analiza el código consolidado retornado por la herramienta. Realizarás un análisis exhaustivo celda por celda y script por script para identificar patrones de código que se correlacionen con las fases de CRISP-DM.
3. Redacta la documentación de forma muy completa, extensa, analítica y técnica en formato Markdown, siguiendo rigurosamente la estructura CRISP-DM detallada abajo.
4. Una vez concluida la redacción completa, invoca de forma autónoma la herramienta 'save_markdown_as_docx' para exportar tu texto Markdown a un archivo Word (.docx).
5. Notifica al usuario en la terminal que el archivo está guardado y ofrécele un breve resumen de los principales hitos del modelo detectados.

REGLAS DE GENERACIÓN OBLIGATORIAS:
- Tono y Estilo: Altamente formal, objetivo, técnico y de nivel Senior (apropiado para directores de Data y Analítica).
- CERO ALUCINACIONES: Basarás tu documentación ÚNICAMENTE en lo que demuestre el código analizado. Si el repositorio carece de elementos de alguna fase, indícalo formalmente: "No se evidencia información de esta fase en el código analizado."

ESTRUCTURA RIGUROSA DE REDACCIÓN (CRISP-DM):

## 1. Comprensión del Negocio (Business Understanding)
- Analiza la introducción del proyecto, comentarios de cabecera y nombres de variables objetivo (targets) para deducir de manera ejecutiva el problema de negocio que se quiere resolver.
- Propón dos KPIs técnicos de Machine Learning y su respectivo impacto estimado en KPIs de negocio de la compañía (p. ej., optimización de costos o retención de clientes).

## 2. Comprensión de los Datos (Data Understanding)
- Identifica y enlista cada fuente de datos que se cargue en el código (queries SQL, archivos csv, parquet, lecturas de APIs).
- Detalla el Análisis Exploratorio de Datos (EDA) realizado: describe qué estadísticas se computaron (p. ej., correlaciones, distribuciones, conteo de nulos) y qué gráficos se generaron en el código para comprender los datos.

## 3. Preparación de los Datos (Data Preparation)
- Explica minuciosamente todas las transformaciones de datos aplicadas en los scripts: imputación de nulos, tratamiento de outliers, transformaciones categóricas (One-Hot Encoding, Label Encoding), escalamiento de datos (MinMaxScaler, StandardScaler) y creación de nuevas variables (Feature Engineering).
- Enlista las librerías utilizadas (Pandas, NumPy, Scipy, etc.).

## 4. Modelado (Modeling)
- Identifica todos los algoritmos que fueron entrenados en el repositorio (ej. XGBoost, regresiones, redes neuronales).
- Crea una tabla en Markdown detallada que documente los hiperparámetros que el científico de datos configuró en el código para inicializar cada modelo.
- Describe la estrategia de validación utilizada (p. ej., Train/Test Split o K-Fold Cross Validation).

## 5. Evaluación (Evaluation)
- Extrae todas las métricas de rendimiento que el código compute (p. ej., RMSE, R2, F1-Score, curva ROC, matrices de confusión).
- Explica técnicamente qué significan estos resultados respecto a la precisión del modelo y su viabilidad para ser usado por el negocio.

## 6. Despliegue (Deployment)
- Identifica las porciones de código orientadas a producción: exportación del modelo (p. ej., pickle, joblib, TensorFlow SavedModel) o integraciones para su puesta en marcha (servicios en la nube, APIs web).
"""

# =====================================================================
# 4. INSTANCIACIÓN DEL AGENTE CON SUS TOOLS
# =====================================================================

root_agent = Agent(
    model='gemini-1.5-pro',
    name='AIDA',
    description="Agente de IA especializado en documentar repositorios completos de ML bajo metodología CRISP-DM utilizando herramientas locales.",
    instruction=CRISP_DM_INSTRUCTIONS,
    tools=[analyze_repository_contents, save_markdown_as_docx]
)
