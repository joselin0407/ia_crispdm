# my_agent_v2/agent.py
import os
import stat
import shutil
import nbformat
import docx
from docx.shared import Pt
from google.adk.agents import Agent

# =====================================================================
# Monitoreo: eliminar solo lectura para errores windows 
# =====================================================================
def remove_readonly(func, path, excinfo):
    """
    Controlador de errores para shutil.rmtree en Windows.
    Quita el atributo de sólo lectura (Read-Only) de los archivos de Git y reintenta.
    """
    os.chmod(path, stat.S_IWRITE)
    func(path)

# =====================================================================
# 1. ANALIZADOR DINÁMICO DE REPOSITORIOS (CON DEPURACIÓN EN TERMINAL)
# =====================================================================

def analyze_repository_contents(repo_path_or_url: str) -> str:
    """
    Escanea un repositorio de Machine Learning estructurado sistemáticamente en las 6 carpetas 
    de CRISP-DM. Puede recibir tanto una ruta de carpeta local como una URL de un repositorio público 
    de GitHub (ej: https://github.com/usuario/proyecto.git).
    
    Args:
        repo_path_or_url: Ruta de la carpeta local O URL de GitHub del repositorio.
    """
    is_github_url = repo_path_or_url.startswith(("http://", "https://")) and "github.com" in repo_path_or_url
    temp_dir = "temp_cloned_repo"
    
    actual_path = repo_path_or_url
    
    if is_github_url:
        print(f"\n[AIDA (Tool)]: 🌐 Detectada URL de GitHub. Clonando de forma temporal en caliente...", end="\r")
        
        # 🌟 CORREGIDO 1: Borrado blindado con remove_readonly antes de clonar
        if os.path.exists(temp_dir):
            shutil.rmtree(temp_dir, onerror=remove_readonly)
            
        try:
            # Clonado veloz del último commit
            subprocess.run(["git", "clone", "--depth", "1", repo_path_or_url, temp_dir], 
                           check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            actual_path = temp_dir
            print(f"[AIDA (Tool)]: 📥 ¡Clonación exitosa! Analizando estructura de fases...", end="\r")
        except Exception as e:
            return f"Error técnico al intentar clonar el repositorio de GitHub: {str(e)}. Asegúrate de que el repositorio sea público y de tener Git instalado en tu terminal."

    if not os.path.exists(actual_path):
        return f"Error: La ruta o URL '{repo_path_or_url}' no pudo ser resuelta o no existe."
        
    try:
        # 1. Buscar el archivo de contexto en la raíz para extraer el código del proyecto
        project_code = None
        context_content = ""
        
        for file in os.listdir(actual_path):
            if file.endswith("_contexto.txt"):
                project_code = file.replace("_contexto.txt", "") # Ej: "gdp_001"
                filepath = os.path.join(actual_path, file)
                
                try:
                    with open(filepath, 'r', encoding='utf-8') as f:
                        context_content = f.read()
                except UnicodeDecodeError:
                    with open(filepath, 'r', encoding='latin-1') as f:
                        context_content = f.read()
                break
                
        if not project_code:
            return "Error: No se encontró ningún archivo de contexto en la raíz (ej: 'gdp_001_contexto.txt') para identificar el código del proyecto."
            
        consolidated_payload = []
        consolidated_payload.append(f"==================================================\n")
        consolidated_payload.append(f"CÓDIGO DE PROYECTO DETECTADO: {project_code}\n")
        consolidated_payload.append(f"CONTEXTO COMERCIAL DE LA COMPAÑÍA:\n{context_content}\n")
        consolidated_payload.append(f"==================================================\n\n")

        # 2. Mapear las 6 carpetas oficiales de la metodología CRISP-DM
        folders_mapping = {
            "1. BU": "COMPRENSIÓN DEL NEGOCIO (Business Understanding)",
            "2. DU": "COMPRENSIÓN DE LOS DATOS (Data Understanding)",
            "3. DP": "PREPARACIÓN DE LOS DATOS (Data Preparation)",
            "4. MOD": "MODELADO (Modeling)",
            "5. EVA": "EVALUACIÓN (Evaluation)",
            "6. DEPL": "DESPLIEGUE (Deployment)"
        }

        files_read_count = 0

        for folder_name, phase_title in folders_mapping.items():
            folder_directory = os.path.join(actual_path, folder_name)
            
            if not os.path.exists(folder_directory):
                consolidated_payload.append(f"--- FASE: {phase_title} ---\n[Sin información: Carpeta '{folder_name}' no encontrada en el repositorio]\n\n")
                continue
                
            phase_files = [f for f in os.listdir(folder_directory) if f.startswith(project_code)]
            
            if not phase_files:
                consolidated_payload.append(f"--- FASE: {phase_title} ---\n[No se encontraron notebooks o códigos para el proyecto '{project_code}' en esta fase]\n\n")
                continue
                
            consolidated_payload.append(f"--- FASE: {phase_title} ---\n")
            
            for file in phase_files:
                filepath = os.path.join(folder_directory, file)
                consolidated_payload.append(f"-> Archivo de la fase: {file}\n")
                files_read_count += 1
                
                # Procesamiento de Scripts Python (.py)
                if file.endswith('.py'):
                    try:
                        try:
                            with open(filepath, 'r', encoding='utf-8') as f:
                                content = f.read()
                        except UnicodeDecodeError:
                            with open(filepath, 'r', encoding='latin-1') as f:
                                content = f.read()
                        consolidated_code_block = f"[CÓDIGO PYTHON ENCONTRADO]:\n{content}\n"
                    except Exception as e:
                        consolidated_code_block = f"[Error leyendo .py]: {str(e)}\n"
                    consolidated_payload.append(consolidated_code_block)
                
                # Procesamiento de Jupyter Notebooks (.ipynb)
                elif file.endswith('.ipynb'):
                    try:
                        with open(filepath, 'r', encoding='utf-8') as f:
                            nb = nbformat.read(f, as_version=4)
                        
                        notebook_text = []
                        for idx, cell in enumerate(nb.cells):
                            if cell.cell_type == 'markdown':
                                notebook_text.append(f"[Comentarios Markdown Celda {idx}]:\n{cell.source}")
                            elif cell.cell_type == 'code':
                                notebook_text.append(f"[Bloque de Código Celda {idx}]:\n{cell.source}")
                                
                        consolidated_payload.append("\n".join(notebook_text) + "\n")
                    except Exception as e:
                        consolidated_payload.append(f"[Error leyendo .ipynb]: {str(e)}\n")
            
            consolidated_payload.append("\n" + "="*40 + "\n\n")

        # DEPURACIÓN DE CONSOLA: Mostramos cuántos archivos se leyeron físicamente
        print(" "*60, end="\r") # Limpia línea anterior
        print(f"[AIDA (Tool)]: 🔍 Escaneo de GitHub completo.")
        print(f"               • Proyecto detectado: '{project_code}'")
        print(f"               • Archivos técnicos leídos de forma exitosa: {files_read_count}")

        # 🌟 CORREGIDO 2: Borrado blindado con remove_readonly en flujo exitoso
        if is_github_url and os.path.exists(temp_dir):
            shutil.rmtree(temp_dir, onerror=remove_readonly)

        return "".join(consolidated_payload)
        
    except Exception as e:
        # 🌟 CORREGIDO 3: Borrado blindado con remove_readonly en bloque de excepción
        if is_github_url and os.path.exists(temp_dir):
            shutil.rmtree(temp_dir, onerror=remove_readonly)
        return f"Error técnico durante el escaneo estructurado del repositorio: {str(e)}"


# =====================================================================
# 2. HERRAMIENTA DE EXPORTACIÓN A WORD (DOCX)
# =====================================================================

def save_markdown_as_docx(markdown_content: str, filename: str = "documento_CRISP_DM.docx") -> str:
    """
    Convierte un texto con formato Markdown en un archivo de Microsoft Word (.docx) formal.
    """
    try:
        doc = docx.Document()
        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)

        doc.add_heading('Documentación Oficial de Proyecto ML (CRISP-DM)', level=0)
        
        lines = markdown_content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            if line.startswith('### '):
                doc.add_heading(line.replace('### ', ''), level=3)
            elif line.startswith('## '):
                doc.add_heading(line.replace('## ', ''), level=2)
            elif line.startswith('# '):
                doc.add_heading(line.replace('# ', ''), level=1)
            elif line.startswith('- '):
                doc.add_paragraph(line.replace('- ', ''), style='List Bullet')
            else:
                doc.add_paragraph(line)
                
        doc.save(filename)
        return f"¡Éxito! El documento de Word ha sido autogenerado y guardado como '{filename}' en tu carpeta de proyecto."
        
    except Exception as e:
        return f"Error técnico al generar el documento Word: {str(e)}"


# =====================================================================
# 3. PROMPT OPERATIVO BLINDADO CONTRA ALUCINACIONES (ESTÁNDAR IBM)
# =====================================================================

CRISP_DM_IBM_INSTRUCTIONS = """
Eres AIDA (Agente Inteligente de Documentación Analítica), un experto técnico senior en MLOps, Gobierno de Datos y la metodología CRISP-DM basada estrictamente en la guía oficial de IBM SPSS Modeler.

Tu objetivo es leer un reporte de código estructurado por carpetas cronológicas (fases), extraer los hallazgos y redactar un documento de gobierno completo y formal.

INSTRUCCIONES DE OPERACIÓN:
1. Si el usuario te indica la carpeta o repositorio, ejecuta inmediatamente la herramienta 'analyze_repository_contents'.
2. Lee el archivo de contexto inyectado (ej: 'gdp_001_contexto.txt') para asimilar la situación comercial, objetivos de negocio y restricciones.
3. Analiza el código de cada una de las 6 carpetas correspondientes a las fases.
4. Genera el documento Word invocando de forma autónoma la herramienta 'save_markdown_as_docx'.
5. Ofrécele al usuario en la terminal una confirmación formal y un resumen ejecutivo estructurado de los hallazgos.

REGLAS DE GENERACIÓN OBLIGATORIAS (MÁXIMA SEGURIDAD):
- Cero Alucinaciones (CRÍTICO): Limítate a documentar EXCLUSIVAMENTE lo que demuestre el código y el archivo de contexto del repositorio inyectado. 
- PROHIBIDO INVENTAR: No utilices los ejemplos del prompt (como la retención de clientes o churn) como si fueran los datos reales si el repositorio es de otra temática. Si el repositorio habla de Iris / BioData, documentarás ÚNICAMENTE sobre Iris / BioData.
- RECHAZO DE REPOSITORIO VACÍO: Si la herramienta te devuelve un error o indica que leyó 0 archivos técnicos, debes detener la ejecución, no generar ningún Word de fantasía, y responder al usuario informando que el repositorio está vacío o mal estructurado para el código del proyecto.

ESTRUCTURA RIGUROSA DE DOCUMENTACIÓN (ALINEADA A IBM SPSS MODELER):

## 1. Comprensión del Negocio (Business Understanding)
- **Determinación de los objetivos comerciales:** Define con claridad el problema comercial real expresado en el archivo de contexto.
- **Evaluación de la situación:** Detalle los recursos identificados, requisitos, supuestos, restricciones y la evaluación de riesgos del proyecto según el contexto de la empresa.
- **Determinación de los objetivos de minería de datos:** Traduzca los objetivos de negocio a objetivos analíticos técnicos de Machine Learning.

## 2. Comprensión de los Datos (Data Understanding)
- **Recopilación de datos iniciales:** Enlista las fuentes de datos reales leídas en el código (BigQuery, GS bucket, CSVs locales).
- **Descripción de los datos:** Detalla el volumen detectado (shape, número de campos, tipo de variables).
- **Exploración de los datos:** Describe qué pruebas estadísticas o de visualización se aprecian en el código (correlaciones, distribuciones, histogramas).
- **Verificación de la calidad de los datos:** Escriba un informe de calidad de datos detectando valores nulos, registros atípicos (anomalías) o conflictos en los datos en el código.

## 3. Preparación de los Datos (Data Preparation)
- **Selección de datos:** Explique qué filtros o subconjuntos de datos se extrajeron para el modelado.
- **Limpieza de datos:** Detalle el tratamiento de valores nulos y tratamiento de duplicados.
- **Construcción e Integración de datos:** Explique la creación de nuevas variables (Feature Engineering) y fusiones (merges/joins) de tablas en el código.

## 4. Modelado (Modeling)
- **Selección de la técnica de modelado:** Argumente qué algoritmo se escogió para entrenar en base al problema.
- **Generación de un plan de prueba:** Detalle la estrategia de partición de datos (Train/Test Split o K-Fold Cross Validation).
- **Construcción y Valoración del modelo:** Liste en una tabla ordenada en Markdown los hiperparámetros de entrenamiento del modelo y describa el comportamiento del entrenamiento.

## 5. Evaluación (Evaluation)
- **Evaluación de resultados:** Reporte las métricas estadísticas obtenidas en el código de validación y contraste si se alcanzan los objetivos comerciales de éxito iniciales.
- **Proceso de revisión:** Realice una reflexión técnica (Lecciones aprendidas, aciertos, errores).
- **Determinación de los pasos siguientes:** Decida formalmente si el modelo está listo para el despliegue o si requiere volver a la fase de modelado para refinamiento.

## 6. Despliegue (Deployment)
- **Planificación de despliegue, supervisión y mantenimiento:** Describa el plan para incorporar el modelo a producción.
"""

# =====================================================================
# 4. INSTANCIACIÓN DEL AGENTE V2
# =====================================================================

root_agent = Agent(
    model='gemini-3.1-flash-lite',
    name='AIDA_V2',
    description="Agente de IA avanzado especializado en documentar repositorios de Machine Learning bajo metodología IBM CRISP-DM estructurado por fases.",
    instruction= CRISP_DM_IBM_INSTRUCTIONS,
    tools=[analyze_repository_contents, save_markdown_as_docx]
)

