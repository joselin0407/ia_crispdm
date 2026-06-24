# my_agent_v3/agent.py
import os
import stat
import shutil
import nbformat
import docx
from docx.shared import Pt
from google.adk.agents import Agent

# =====================================================================
# AUXILIAR: CERRAJERO DE WINDOWS (ELIMINA CANDADOS DE SOLO LECTURA)
# =====================================================================
def remove_readonly(func, path, excinfo):
    """
    Controlador de errores para shutil.rmtree en Windows.
    Quita el atributo de sólo lectura (Read-Only) de los archivos de Git y reintenta.
    """
    os.chmod(path, stat.S_IWRITE)
    func(path)

# Helper para parsear URIs de Google Cloud Storage
def parse_gcs_uri(gcs_uri: str):
    """
    Parsea una URI de GCS (gs://bucket_name/path/to/folder) y retorna (bucket_name, prefix).
    """
    if gcs_uri.startswith("gs://"):
        gcs_uri = gcs_uri[5:]
    parts = gcs_uri.split("/", 1)
    bucket_name = parts[0]
    prefix = parts[1] if len(parts) > 1 else ""
    if prefix and not prefix.endswith("/"):
        prefix += "/"
    return bucket_name, prefix

# =====================================================================
# 1. ANALIZADOR HÍBRIDO DE REPOSITORIOS (SOPORTA LOCAL, GITHUB Y GCS BUCKET)
# =====================================================================

def analyze_repository_contents(repo_path_or_url: str) -> str:
    """
    Escanea un repositorio de Machine Learning estructurado sistemáticamente en las 6 carpetas 
    de CRISP-DM. Soporta rutas locales de carpetas, URLs de GitHub y URIs de Google Cloud Storage 
    (ej: gs://mi-bucket-vanti/gdp_001).
    
    Args:
        repo_path_or_url: Ruta local, URL de GitHub o URI gs:// de Google Cloud Storage.
    """
    import os
    import nbformat
    import json
    
    # Identificar el protocolo de entrada
    is_github_url = repo_path_or_url.startswith(("http://", "https://")) and "github.com" in repo_path_or_url
    is_gcs = repo_path_or_url.startswith("gs://")
    temp_dir = "temp_cloned_repo"
    
    # -----------------------------------------------------------------
    # MODO 1: LECTURA DESDE GOOGLE CLOUD STORAGE (GCS BUCKET)
    # -----------------------------------------------------------------
    if is_gcs:
        try:
            from google.cloud import storage
        except ImportError:
            return "Error: Para leer desde Google Cloud Storage, debes instalar la librería oficial ejecutando: pip install google-cloud-storage"
            
        bucket_name, base_prefix = parse_gcs_uri(repo_path_or_url)
        print(f"\n[AIDA (Tool)]: ☁️ Conectando a Google Cloud Storage Bucket: '{bucket_name}'...", end="\r")
        
        try:
            # Inicializa el cliente de GCS. Buscará las credenciales locales de gcloud (ADC) automáticamente.
            storage_client = storage.Client()
            bucket = storage_client.bucket(bucket_name)
            
            # Buscar el archivo de contexto en la raíz del prefijo de GCS
            project_code = None
            context_content = ""
            
            # Listamos los elementos del primer nivel en GCS
            blobs = storage_client.list_blobs(bucket_name, prefix=base_prefix, delimiter="/")
            for blob in blobs:
                filename = os.path.basename(blob.name)
                if filename.endswith("_contexto.txt"):
                    project_code = filename.replace("_contexto.txt", "") # Ej: "gdp_001"
                    try:
                        context_content = blob.download_as_text(encoding="utf-8")
                    except Exception:
                        context_content = blob.download_as_text(encoding="latin-1")
                    break
                    
            if not project_code:
                return f"Error: No se encontró ningún archivo de contexto terminado en '_contexto.txt' en la ruta de GCS: {repo_path_or_url}"
                
            consolidated_payload = []
            consolidated_payload.append(f"CÓDIGO DE PROYECTO DETECTADO: {project_code}\n")
            consolidated_payload.append(f"CONTEXTO COMERCIAL DE LA COMPAÑÍA:\n{context_content}\n")
            consolidated_payload.append(f"==================================================\n\n")

            folders_mapping = {
                "1. BU": "COMPRENSIÓN DEL NEGOCIO (Business Understanding)",
                "2. DU": "COMPRENSIÓN DE LOS DATOS (Data Understanding)",
                "3. DP": "PREPARACIÓN DE LOS DATOS (Data Preparation)",
                "4. MOD": "MODELADO (Modeling)",
                "5. EVA": "EVALUACIÓN (Evaluation)",
                "6. DEPL": "DESPLIEGUE (Deployment)"
            }

            files_read_count = 0

            # Recorrer las fases en GCS
            for folder_name, phase_title in folders_mapping.items():
                folder_prefix = f"{base_prefix}{folder_name}/"
                phase_blobs = list(storage_client.list_blobs(bucket_name, prefix=folder_prefix))
                
                # Filtrar archivos del proyecto actual
                valid_blobs = [b for b in phase_blobs if os.path.basename(b.name).startswith(project_code)]
                
                if not valid_blobs:
                    consolidated_payload.append(f"--- FASE: {phase_title} ---\n[No se encontraron notebooks o códigos para el proyecto '{project_code}' en esta fase en GCS]\n\n")
                    continue
                    
                consolidated_payload.append(f"--- FASE: {phase_title} ---\n")
                for blob in valid_blobs:
                    filename = os.path.basename(blob.name)
                    consolidated_payload.append(f"-> Archivo en GCS: {filename}\n")
                    files_read_count += 1
                    
                    if filename.endswith(".py"):
                        try:
                            try:
                                content = blob.download_as_text(encoding="utf-8")
                            except Exception:
                                content = blob.download_as_text(encoding="latin-1")
                            consolidated_payload.append(f"[CÓDIGO PYTHON ENCONTRADO]:\n{content}\n")
                        except Exception as e:
                            consolidated_payload.append(f"[Error leyendo .py desde GCS]: {str(e)}\n")
                    elif filename.endswith(".ipynb"):
                        try:
                            notebook_bytes = blob.download_as_bytes()
                            nb = nbformat.reads(notebook_bytes.decode("utf-8"), as_version=4)
                            notebook_text = []
                            for idx, cell in enumerate(nb.cells):
                                if cell.cell_type == 'markdown':
                                    notebook_text.append(f"[Comentarios Markdown Celda {idx}]:\n{cell.source}")
                                elif cell.cell_type == 'code':
                                    notebook_text.append(f"[Bloque de Código Celda {idx}]:\n{cell.source}")
                            consolidated_payload.append("\n".join(notebook_text) + "\n")
                        except Exception as e:
                            consolidated_payload.append(f"[Error leyendo .ipynb desde GCS]: {str(e)}\n")
                
                consolidated_payload.append("\n" + "="*40 + "\n\n")

            print(" "*60, end="\r")
            print(f"[AIDA (Tool)]: ☁️ Escaneo de Google Cloud Storage completo.")
            print(f"               • Proyecto detectado: '{project_code}'")
            print(f"               • Archivos leídos desde GCS: {files_read_count}")
            
            return "".join(consolidated_payload)
            
        except Exception as e:
            return f"Error de conexión con GCS: {str(e)}. Valida que estés logueado en gcloud en tu máquina."

    # -----------------------------------------------------------------
    # MODO 2: CLONADO TEMPORAL DESDE GITHUB (LINK WEB)
    # -----------------------------------------------------------------
    actual_path = repo_path_or_url
    if is_github_url:
        print(f"\n[AIDA (Tool)]: 🌐 Detectada URL de GitHub. Clonando de forma temporal en caliente...", end="\r")
        if os.path.exists(temp_dir):
            shutil.rmtree(temp_dir, onerror=remove_readonly)
        try:
            import subprocess
            subprocess.run(["git", "clone", "--depth", "1", repo_path_or_url, temp_dir], 
                           check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            actual_path = temp_dir
            print(f"[AIDA (Tool)]: 📥 ¡Clonación exitosa! Analizando estructura de fases...", end="\r")
        except Exception as e:
            return f"Error técnico al intentar clonar el repositorio de GitHub: {str(e)}."

    # -----------------------------------------------------------------
    # MODO 3: ESCANEO DE CARPETA LOCAL (VS CODE)
    # -----------------------------------------------------------------
    if not os.path.exists(actual_path):
        return f"Error: La ruta o URL '{repo_path_or_url}' no pudo ser resuelta o no existe."
        
    try:
        project_code = None
        context_content = ""
        
        for file in os.listdir(actual_path):
            if file.endswith("_contexto.txt"):
                project_code = file.replace("_contexto.txt", "")
                filepath = os.path.join(actual_path, file)
                try:
                    with open(filepath, 'r', encoding='utf-8') as f:
                        context_content = f.read()
                except UnicodeDecodeError:
                    with open(filepath, 'r', encoding='latin-1') as f:
                        context_content = f.read()
                break
                
        if not project_code:
            return "Error: No se encontró ningún archivo de contexto en la raíz (ej: 'gdp_001_contexto.txt') para identificar el código del proyecto."
            
        consolidated_payload = []
        consolidated_payload.append(f"CÓDIGO DE PROYECTO DETECTADO: {project_code}\n")
        consolidated_payload.append(f"CONTEXTO COMERCIAL DE LA COMPAÑÍA:\n{context_content}\n")
        consolidated_payload.append(f"==================================================\n\n")

        folders_mapping = {
            "1. BU": "COMPRENSIÓN DEL NEGOCIO (Business Understanding)",
            "2. DU": "COMPRENSIÓN DE LOS DATOS (Data Understanding)",
            "3. DP": "PREPARACIÓN DE LOS DATOS (Data Preparation)",
            "4. MOD": "MODELADO (Modeling)",
            "5. EVA": "EVALUACIÓN (Evaluation)",
            "6. DEPL": "DESPLIEGUE (Deployment)"
        }

        files_read_count = 0
        for folder_name, phase_title in folders_mapping.items():
            folder_directory = os.path.join(actual_path, folder_name)
            if not os.path.exists(folder_directory):
                consolidated_payload.append(f"--- FASE: {phase_title} ---\n[Sin información: Carpeta '{folder_name}' no encontrada]\n\n")
                continue
                
            phase_files = [f for f in os.listdir(folder_directory) if f.startswith(project_code)]
            if not phase_files:
                consolidated_payload.append(f"--- FASE: {phase_title} ---\n[No se encontraron archivos para '{project_code}' en esta fase]\n\n")
                continue
                
            consolidated_payload.append(f"--- FASE: {phase_title} ---\n")
            for file in phase_files:
                filepath = os.path.join(folder_directory, file)
                consolidated_payload.append(f"-> Archivo local: {file}\n")
                files_read_count += 1
                
                if file.endswith('.py'):
                    try:
                        try:
                            with open(filepath, 'r', encoding='utf-8') as f:
                                content = f.read()
                        except UnicodeDecodeError:
                            with open(filepath, 'r', encoding='latin-1') as f:
                                content = f.read()
                        consolidated_payload.append(f"[CÓDIGO PYTHON ENCONTRADO]:\n{content}\n")
                    except Exception as e:
                        consolidated_payload.append(f"[Error leyendo .py]: {str(e)}\n")
                elif file.endswith('.ipynb'):
                    try:
                        with open(filepath, 'r', encoding='utf-8') as f:
                            nb = nbformat.read(f, as_version=4)
                        notebook_text = []
                        for idx, cell in enumerate(nb.cells):
                            if cell.cell_type == 'markdown':
                                notebook_text.append(f"[Comentarios Markdown Celda {idx}]:\n{cell.source}")
                            elif cell.cell_type == 'code':
                                notebook_text.append(f"[Bloque de Código Celda {idx}]:\n{cell.source}")
                        consolidated_payload.append("\n".join(notebook_text) + "\n")
                    except Exception as e:
                        consolidated_payload.append(f"[Error leyendo .ipynb]: {str(e)}\n")
            
            consolidated_payload.append("\n" + "="*40 + "\n\n")

        print(" "*60, end="\r")
        print(f"[AIDA (Tool)]: 🔍 Escaneo de carpeta local completo.")
        print(f"               • Proyecto detectado: '{project_code}'")
        print(f"               • Archivos leídos localmente: {files_read_count}")

        if is_github_url and os.path.exists(temp_dir):
            shutil.rmtree(temp_dir, onerror=remove_readonly)
        return "".join(consolidated_payload)
        
    except Exception as e:
        if is_github_url and os.path.exists(temp_dir):
            shutil.rmtree(temp_dir, onerror=remove_readonly)
        return f"Error técnico durante el escaneo local: {str(e)}"


# =====================================================================
# 2. HERRAMIENTA DE EXPORTACIÓN A WORD (DOCX) - ROBUSTA Y ULTRA LIMPIA
# =====================================================================

def save_markdown_as_docx(markdown_content: str, filename: str = "documento_CRISP_DM.docx") -> str:
    """
    Convierte texto Markdown en Microsoft Word (.docx) formal. 
    Limpia los caracteres especiales visuales de consola y asteriscos de formato Markdown.
    """
    try:
        doc = docx.Document()
        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)
        doc.add_heading('Documentación Oficial de Proyecto ML (CRISP-DM)', level=0)
        
        lines = markdown_content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # Omitimos líneas de separación visual de consola
            if line.startswith("===") or line.startswith("---") or line.startswith("═══") or line.startswith("───"):
                continue
                
            # Formateo de Encabezados (Limpia asteriscos de negrita si existen)
            if line.startswith('### ') or line.startswith('## ') or line.startswith('# '):
                level = 3 if line.startswith('### ') else (2 if line.startswith('## ') else 1)
                header_text = line.replace('### ', '').replace('## ', '').replace('# ', '')
                header_text = header_text.replace('**', '').replace('*', '').strip()
                doc.add_heading(header_text, level=level)
                
            # Formateo de Viñetas (Soporta - , * , + y viñetas unificadas · )
            elif line.startswith(('- ', '* ', '+ ', '· ')):
                bullet_text = line[2:].replace('**', '').replace('*', '').strip()
                doc.add_paragraph(bullet_text, style='List Bullet')
                
            else:
                # Líneas normales (Limpia asteriscos huérfanos para estética profesional)
                clean_line = line.replace('**', '').replace('*', '')
                doc.add_paragraph(clean_line)
                
        doc.save(filename)
        return f"¡Éxito! El documento de Word ha sido autogenerado y guardado como '{filename}' en tu carpeta de proyecto."
        
    except Exception as e:
        return f"Error técnico al generar el documento Word: {str(e)}"


# =====================================================================
# 3. PROMPT OPERATIVO BLINDADO CONTRA ALUCINACIONES (ESTÁNDAR IBM)
# =====================================================================

CRISP_DM_IBM_INSTRUCTIONS = """
Eres AIDA (Agente Inteligente de Documentación Analítica), un experto técnico senior en MLOps, Gobierno de Datos y la metodología CRISP-DM basada estrictamente en la guía oficial de IBM SPSS Modeler.

Tu objetivo es leer un reporte de código estructurado por carpetas cronológicas (fases), extraer los hallazgos y redactar un documento de gobierno completo y formal.

INSTRUCCIONES DE OPERACIÓN:
1. Si el usuario te indica la carpeta, la URL de GitHub o la URI de Google Cloud Storage (gs://), ejecuta inmediatamente la herramienta 'analyze_repository_contents'.
2. Analiza el código de cada una de las 6 carpetas correspondientes a las fases y tradúcelos al reporte formal bajo la estructura rigurosa de IBM SPSS Modeler detallada abajo.
3. Genera el documento Word invocando de forma autónoma la herramienta 'save_markdown_as_docx'.
4. Ofrécele al usuario en la terminal una confirmación formal y un resumen ejecutivo estructurado de los hallazgos.

REGLAS DE GENERACIÓN OBLIGATORIAS (MÁXIMA SEGURIDAD):
- Cero Alucinaciones (CRÍTICO): Limítate a documentar EXCLUSIVAMENTE lo que demuestre el código y el archivo de contexto del repositorio inyectado. 
- PROHIBIDO INVENTAR: No utilices ejemplos de retención (churn) si el repositorio habla de otro tema. Si el repositorio habla de Iris / BioData, documentarás ÚNICAMENTE sobre Iris / BioData.
- RECHAZO DE REPOSITORIO VACÍO: Si la herramienta te devuelve un error o indica que leyó 0 archivos técnicos, debes detener la ejecución, no generar ningún Word, y responder al usuario informando sobre el fallo.

ESTRUCTURA RIGUROSA DE DOCUMENTACIÓN (ALINEADA A IBM SPSS MODELER):

## 1. Comprensión del Negocio (Business Understanding)
- **Determinación de los objetivos comerciales:** Define con claridad el problema comercial real expresado en el archivo de contexto.
- **Evaluación de la situación:** Detalle los recursos identificados, requisitos, supuestos, restricciones y la evaluación de riesgos del proyecto según el contexto de la empresa.
- **Determinación de los objetivos de minería de datos:** Traduzca los objetivos de negocio a objetivos analíticos técnicos de Machine Learning.

## 2. Comprensión de los Datos (Data Understanding)
- **Recopilación de datos iniciales:** Enlista las fuentes de datos reales leídas en el código (BigQuery, GS bucket, CSVs locales).
- **Descripción de los datos:** Detalla el volumen detectado (shape, número de campos, tipo de variables).
- **Exploración de los datos:** Describe qué pruebas estadísticas o de visualización se aprecian en el código (correlaciones, distribuciones, histogramas).
- **Verificación de la calidad de los datos:** Escriba un informe de calidad de datos detectando valores nulos, registros atípicos (anomalías) o de mala calidad.

## 3. Preparación de los Datos (Data Preparation)
- **Selección de datos:** Explique qué filtros o subconjuntos de datos se extrajeron para el modelado.
- **Limpieza de datos:** Detalle el tratamiento de valores nulos y tratamiento de duplicados.
- **Construcción e Integración de datos:** Explique la creación de nuevas variables (Feature Engineering) y fusiones (merges/joins) de tablas en el código.

## 4. Modelado (Modeling)
- **Selección de la técnica de modelado:** Argumente qué algoritmo se escogió para entrenar en base al problema.
- **Generación de un plan de prueba:** Detalle la estrategia de partición de datos (Train/Test Split o K-Fold Cross Validation).
- **Construcción y Valoración del modelo:** Liste en una tabla ordenada en Markdown los hiperparámetros de entrenamiento del modelo y describa el comportamiento del entrenamiento.

## 5. Evaluación (Evaluation)
- **Evaluación de resultados:** Reporte las métricas estadísticas obtenidas en el código de validación y contraste si se alcanzan los objetivos comerciales de éxito iniciales.
- **Proceso de revisión:** Realice una reflexión técnica (Lecciones aprendidas, aciertos, errores).
- **Determinación de los pasos siguientes:** Decida formalmente si el modelo está listo para el despliegue o si requiere volver a la fase de modelado para refinamiento.

## 6. Despliegue (Deployment)
- **Planificación de despliegue, supervisión y mantenimiento:** Describa el plan para incorporar el modelo a producción.
"""

# =====================================================================
# 4. INSTANCIACIÓN DEL AGENTE V3
# =====================================================================

root_agent = Agent(
    model='gemini-3.1-flash-lite',
    name='AIDA_V3',
    description="Agente de IA avanzado especializado en documentar repositorios de Machine Learning bajo metodología IBM CRISP-DM estructurado por fases.",
    instruction=CRISP_DM_IBM_INSTRUCTIONS,
    tools=[analyze_repository_contents, save_markdown_as_docx]
)

